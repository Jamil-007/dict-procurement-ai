import io
from pathlib import Path
from docxtpl import DocxTemplate
from docx import Document

# Lighter draft footer for Group A (procuring-entity authored) forms.
GROUP_A_DISCLAIMER = "Generated draft — verify before submission."

# Fixed-position market tables. The Market Scoping form has a fixed activities table
# (6 predefined activities + an "Other" row) and a fixed results table (6 parameters).
# We map canonical keys -> table row index (row 0 is the header row).
_MARKET_ACTIVITY_TABLE = 2
_MARKET_RESULT_TABLE = 3
_MARKET_ACTIVITY_ROWS = {
    "consultation": 1,
    "summit": 2,
    "technical_review": 3,
    "product_review": 4,
    "price_sourcing": 5,
    "philgeps": 6,
}
_MARKET_RESULT_ROWS = {
    "cost_estimate": 1,
    "design_spec": 2,
    "technical_criteria": 3,
    "delivery": 4,
    "storage": 5,
    "risk": 6,
}
_CHECKED_GLYPH = "☑"  # ☑


def render_docx(template_path: Path, context: dict, disclaimer: str | None = None) -> bytes:
    tpl = DocxTemplate(str(template_path))
    ctx = {k: ("[TBD]" if v is None else v) for k, v in context.items()}
    if disclaimer is not None:
        ctx["disclaimer"] = disclaimer
    tpl.render(ctx)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def _attr(obj, name):
    """Read an attribute whether obj is a pydantic model or a plain dict."""
    if obj is None:
        return None
    if hasattr(obj, name):
        return getattr(obj, name)
    if isinstance(obj, dict):
        return obj.get(name)
    return None


def _set_cell_text(cell, text: str) -> None:
    cell.text = text


def fill_market(template_path: Path, data) -> bytes:
    """Render the Market Scoping form: header tokens via docxtpl, then write the fixed
    activities/results tables directly with python-docx. `data` is a MarketScopingData."""
    ctx = data.model_dump()
    # Composite fields are written to tables below, not to docxtpl placeholders.
    activity_flags = ctx.pop("activity_flags", {}) or {}
    result_rows = ctx.pop("result_rows", {}) or {}
    rendered = render_docx(template_path, ctx, disclaimer=GROUP_A_DISCLAIMER)

    doc = Document(io.BytesIO(rendered))
    tables = doc.tables
    if len(tables) > _MARKET_ACTIVITY_TABLE:
        act = tables[_MARKET_ACTIVITY_TABLE]
        for key, row_idx in _MARKET_ACTIVITY_ROWS.items():
            entry = activity_flags.get(key) if isinstance(activity_flags, dict) else None
            if entry is not None and _attr(entry, "checked") and row_idx < len(act.rows):
                _set_cell_text(act.rows[row_idx].cells[0], _CHECKED_GLYPH)
                doc_text = _attr(entry, "documentation")
                if doc_text and len(act.rows[row_idx].cells) > 2:
                    _set_cell_text(act.rows[row_idx].cells[2], str(doc_text))
    if len(tables) > _MARKET_RESULT_TABLE:
        res = tables[_MARKET_RESULT_TABLE]
        for key, row_idx in _MARKET_RESULT_ROWS.items():
            entry = result_rows.get(key) if isinstance(result_rows, dict) else None
            if entry is not None and row_idx < len(res.rows):
                considered = _attr(entry, "considered")
                recommendation = _attr(entry, "recommendation")
                cells = res.rows[row_idx].cells
                if considered and len(cells) > 1:
                    _set_cell_text(cells[1], str(considered))
                if recommendation and len(cells) > 2:
                    _set_cell_text(cells[2], str(recommendation))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
