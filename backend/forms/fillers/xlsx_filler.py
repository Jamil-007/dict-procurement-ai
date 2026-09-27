import io
from pathlib import Path
import openpyxl
from forms.schemas import PPMPData, APPData

def _v(x): return "[TBD]" if x is None else x

def fill_ppmp(template_path: Path, data: PPMPData) -> bytes:
    wb = openpyxl.load_workbook(str(template_path))
    ws = wb["Indicative PPMP (2)"]
    ws["A10"] = f"Fiscal Year : {_v(data.fiscal_year)}"
    ws["A11"] = f"End-User or Implementing Unit: {_v(data.end_user_unit)}"
    r = 18  # first project data row (band 18-22)
    ws.cell(r, 1, _v(data.general_description))
    ws.cell(r, 2, _v(data.project_type))
    ws.cell(r, 3, _v(data.quantity_size))
    ws.cell(r, 4, _v(data.mode_of_procurement))
    ws.cell(r, 5, _v(data.pre_procurement_conference))
    ws.cell(r, 6, _v(data.start_activity))
    ws.cell(r, 7, _v(data.end_activity))
    ws.cell(r, 8, _v(data.expected_delivery))
    ws.cell(r, 9, _v(data.source_of_funds))
    ws.cell(r, 10, _v(data.estimated_budget))
    ws.cell(r, 11, _v(data.supporting_documents))
    ws.cell(r, 12, _v(data.remarks))
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

# Data row for each category section (header row + 1). The "General Requirements"
# section (header at row 10) has its first data row at 11, which in the official
# template holds a sample PAP example (A11 = 10101010) that must be cleared before use.
_APP_SECTION_ROW = {"General Requirements": 11, "Miscellaneous Items": 17, "Common Use Supplies": 23}

# Column indices resolved from the official "APP " sheet header row 7. Do NOT hardcode
# these blindly elsewhere — the header text is the source of truth (see test).
_APP_COL_PROJECT_TITLE = 3   # C7: "Project Title"
_APP_COL_MODE = 7            # G7: "Mode of Procurement"
_APP_COL_BUDGET = 13         # M7: "Estimated Budget / Approved Budget for the Contract (PhP)"

def fill_app(template_path: Path, data: APPData) -> bytes:
    wb = openpyxl.load_workbook(str(template_path))
    ws = wb["APP "]  # trailing space is intentional
    ws["C3"] = f"ANNUAL PROCUREMENT PLAN FOR FY {_v(data.fiscal_year)}"
    ws["C4"] = _v(data.variant)
    row = _APP_SECTION_ROW.get(data.category or "General Requirements", 11)
    # Clear the template's sample PAP example (PAP code / object code) for this row so the
    # generated project does not collide with the sample line. Note: openpyxl treats
    # ws.cell(r, c, None) as a no-op, so assign .value explicitly.
    ws.cell(row, 1).value = None
    ws.cell(row, 2).value = None
    ws.cell(row, _APP_COL_PROJECT_TITLE, _v(data.project_title))
    ws.cell(row, _APP_COL_MODE, _v(data.mode_of_procurement))
    ws.cell(row, _APP_COL_BUDGET, _v(data.estimated_budget))
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
