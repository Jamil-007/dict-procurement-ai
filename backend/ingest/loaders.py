"""The single front door for turning an uploaded file into markdown.

Task T3 accepts any procurement artefact, and several of them (PPMP, APP, the
procurement readiness checklist, RFQ and PBD templates) ship as .docx or
.xlsx rather than PDF. `load_document()` dispatches on extension and, for
PDFs, on whether a usable text layer exists.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, List, Optional

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md"}


@dataclass
class LoadedDocument:
    """A document reduced to markdown, with provenance for the checkers."""

    path: str
    filename: str
    text: str
    total_pages: int = 0
    pages_read: int = 0
    skipped_pages: List[int] = field(default_factory=list)
    source: str = ""  # "text_layer" | "vision" | "mixed" | "docx" | "xlsx" | "plain"
    cache_hits: int = 0
    vision_calls: int = 0
    error: Optional[str] = None

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


def _load_docx(path: Path) -> LoadedDocument:
    from docx import Document

    doc = Document(str(path))
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            parts.append(text)

    for index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        parts.append(f"\n**Table {index}**\n")
        parts.append("| " + " | ".join(rows[0]) + " |")
        parts.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        for row in rows[1:]:
            parts.append("| " + " | ".join(row) + " |")

    return LoadedDocument(
        path=str(path),
        filename=path.name,
        text="\n\n".join(parts),
        source="docx",
    )


def _load_xlsx(path: Path) -> LoadedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    parts: List[str] = []

    for sheet in workbook.worksheets:
        rows = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue

        parts.append(f"## Sheet: {sheet.title}")
        width = max(len(row) for row in rows)
        padded = [row + [""] * (width - len(row)) for row in rows]
        parts.append("| " + " | ".join(padded[0]) + " |")
        parts.append("| " + " | ".join("---" for _ in range(width)) + " |")
        for row in padded[1:]:
            parts.append("| " + " | ".join(row) + " |")

    workbook.close()
    return LoadedDocument(
        path=str(path),
        filename=path.name,
        text="\n\n".join(parts),
        source="xlsx",
    )


def _load_plain(path: Path) -> LoadedDocument:
    return LoadedDocument(
        path=str(path),
        filename=path.name,
        text=path.read_text(encoding="utf-8", errors="replace"),
        source="plain",
    )


def _load_pdf(
    path: Path,
    force_ocr: bool,
    progress: Optional[Callable[[str], None]],
) -> LoadedDocument:
    from ingest.ocr import ocr_pdf

    result = ocr_pdf(path, force=force_ocr, progress=progress)
    sources = {page.source for page in result.pages}
    if sources <= {"text_layer"}:
        source = "text_layer"
    elif sources <= {"vision", "cache"}:
        source = "vision"
    else:
        source = "mixed"

    return LoadedDocument(
        path=str(path),
        filename=path.name,
        text=result.text,
        total_pages=result.total_pages,
        pages_read=len(result.pages),
        skipped_pages=result.skipped_pages,
        source=source,
        cache_hits=result.cache_hits,
        vision_calls=result.vision_calls,
    )


def load_document(
    file_path: str | Path,
    force_ocr: bool = False,
    progress: Optional[Callable[[str], None]] = None,
) -> LoadedDocument:
    """Load any supported document as markdown.

    PDFs with a real text layer are read directly; scanned PDFs go through
    Claude vision. Never raises for an unreadable file -- the failure is
    recorded on `LoadedDocument.error` so one bad upload cannot abort a batch.
    """
    path = Path(file_path)
    if not path.exists():
        return LoadedDocument(
            path=str(path),
            filename=path.name,
            text="",
            error=f"File not found: {path}",
        )

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _load_pdf(path, force_ocr, progress)
        if suffix == ".docx":
            return _load_docx(path)
        if suffix in (".xlsx", ".xlsm"):
            return _load_xlsx(path)
        if suffix in (".txt", ".md"):
            return _load_plain(path)
    except Exception as exc:  # noqa: BLE001 - one bad file must not kill a batch
        return LoadedDocument(
            path=str(path),
            filename=path.name,
            text="",
            error=f"Failed to read {path.name}: {exc}",
        )

    return LoadedDocument(
        path=str(path),
        filename=path.name,
        text="",
        error=(
            f"Unsupported file type '{suffix}'. Supported: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        ),
    )
