from pathlib import Path
from docx import Document
import io
from agents.doc_generation.fillers.docx_filler import render_docx, fill_market
from agents.doc_generation.schemas import MarketScopingData, MarketActivity, MarketResult

FIX = Path(__file__).parent / "fixtures" / "mini.docx"
MARKET_TPL = Path("templates/forms/market.docx")

def test_render_replaces_and_tbd():
    out = render_docx(FIX, {"title": "Hello", "missing": None}, disclaimer="DRAFT")
    doc = Document(io.BytesIO(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Hello" in text and "DRAFT" in text and "[TBD]" in text


def test_fill_market_checks_activity_and_results():
    data = MarketScopingData(
        project_name="GECS Study",
        activity_flags={"consultation": MarketActivity(checked=True)},
        result_rows={"cost_estimate": MarketResult(considered="Yes", recommendation="Aligned")},
    )
    out = fill_market(MARKET_TPL, data)
    doc = Document(io.BytesIO(out))
    act = doc.tables[2]
    assert any("☑" in c.text for row in act.rows for c in row.cells)
    res = doc.tables[3]
    res_text = "\n".join(c.text for row in res.rows for c in row.cells)
    assert "Aligned" in res_text
    # Group A draft footer disclaimer present.
    assert "verify before submission" in "\n".join(p.text for p in doc.paragraphs)
