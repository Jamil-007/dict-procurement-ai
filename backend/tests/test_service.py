from agents.doc_generation import service

def test_generate_ppmp_no_docs_blank(monkeypatch, tmp_path):
    from config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    from utils.storage import generate_thread_id
    tid = generate_thread_id()
    out = service.generate(tid, ["ppmp"], overrides={})
    assert len(out) == 1 and out[0][0].endswith(".xlsx") and out[0][1][:2] == b"PK"

def test_generate_group_b_has_disclaimer(monkeypatch, tmp_path):
    from config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    from utils.storage import generate_thread_id
    import io
    from docx import Document
    tid = generate_thread_id()
    out = service.generate(tid, ["bsd"], overrides={})
    text = "\n".join(p.text for p in Document(io.BytesIO(out[0][1])).paragraphs)
    assert "DRAFT" in text


def test_extract_all_group_b_returns_header_keys(monkeypatch):
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    result = service.extract_all("x", ["bsd", "oss"])
    for key in ["bsd", "oss"]:
        assert set(result[key]["fields"].keys()) == {
            "procuring_entity",
            "project_title",
            "project_reference",
        }
        assert result[key]["group"] == "B_annex"


def test_generate_group_b_stamps_project_title(monkeypatch, tmp_path):
    from config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    from utils.storage import generate_thread_id
    import io
    from docx import Document
    tid = generate_thread_id()
    out = service.generate(
        tid, ["oss"], overrides={"oss": {"project_title": "GECS Laptop Project"}}
    )
    text = "\n".join(p.text for p in Document(io.BytesIO(out[0][1])).paragraphs)
    assert "GECS Laptop Project" in text
    assert "DRAFT" in text


def test_generate_market_with_composites_does_not_raise(monkeypatch, tmp_path):
    from config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    from utils.storage import generate_thread_id
    tid = generate_thread_id()
    overrides = {
        "market": {
            "project_name": "GECS Market Study",
            "activity_flags": {"consultation": {"checked": True}},
            "result_rows": {"cost_estimate": {"considered": "Yes", "recommendation": "Aligned"}},
        }
    }
    out = service.generate(tid, ["market"], overrides=overrides)
    assert len(out) == 1 and out[0][1][:2] == b"PK"


def test_generate_market_malformed_override_no_500(monkeypatch, tmp_path):
    from config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_source_text", lambda tid: "")
    from utils.storage import generate_thread_id
    tid = generate_thread_id()
    # activity_flags as a corrupted string (what "[object Object]" edits would produce).
    out = service.generate(
        tid, ["market"], overrides={"market": {"activity_flags": "[object Object]"}}
    )
    assert len(out) == 1 and out[0][1][:2] == b"PK"


def test_detect_per_file_documents(monkeypatch):
    # Each uploaded PDF is classified separately; documents[] has one entry per file.
    docs = [
        {"filename": "TOR.pdf", "text": "terms of reference body"},
        {"filename": "Market.pdf", "text": "market study body"},
    ]
    monkeypatch.setattr("agents.doc_generation.service.get_source_documents", lambda tid: docs)

    def fake_classify(text):
        if "terms of reference" in text:
            return ["Terms of Reference"]
        if "market study" in text:
            return ["Market Study"]
        return ["Other"]

    monkeypatch.setattr("agents.doc_generation.service.classify_documents", fake_classify)

    out = service.detect("any-thread")

    assert "documents" in out
    assert out["documents"] == [
        {"filename": "TOR.pdf", "doc_types": ["Terms of Reference"]},
        {"filename": "Market.pdf", "doc_types": ["Market Study"]},
    ]
    # doc_types is the ordered union across files (backward compatible key).
    assert out["doc_types"] == ["Terms of Reference", "Market Study"]
    assert "forms" in out and out["forms"]["ppmp"]["recommended"] is True


def test_detect_empty_session(monkeypatch):
    monkeypatch.setattr("agents.doc_generation.service.get_source_documents", lambda tid: [])
    out = service.detect("any-thread")
    assert out["documents"] == []
    assert out["doc_types"] == []
    assert "forms" in out


def test_detect_for_ref_reports_present_and_missing_source_docs(monkeypatch):
    # ppmp needs Terms of Reference or Cost Breakdown; only TOR is on record.
    monkeypatch.setattr(
        "agents.doc_generation.service.recommend_for_ref",
        lambda ref: {
            "doc_types": ["Terms of Reference"],
            "forms": service.recommendations(["Terms of Reference"]),
        },
    )

    out = service.detect_for_ref("PROC-1")

    assert out["doc_types"] == ["Terms of Reference"]
    ppmp = out["forms"]["ppmp"]
    assert ppmp["recommended"] is True
    assert ppmp["source_doc_types"] == ["Terms of Reference", "Cost Breakdown"]
    assert ppmp["present"] == ["Terms of Reference"]
    assert ppmp["missing"] == ["Cost Breakdown"]

    # Group B forms have no source_doc_types, so nothing is ever missing.
    bsd = out["forms"]["bsd"]
    assert bsd["source_doc_types"] == []
    assert bsd["present"] == []
    assert bsd["missing"] == []


def test_detect_for_ref_no_documents_everything_missing(monkeypatch):
    monkeypatch.setattr(
        "agents.doc_generation.service.recommend_for_ref",
        lambda ref: {"doc_types": [], "forms": service.recommendations([])},
    )

    out = service.detect_for_ref("PROC-empty")

    assert out["doc_types"] == []
    assert out["forms"]["ppmp"]["present"] == []
    assert out["forms"]["ppmp"]["missing"] == ["Terms of Reference", "Cost Breakdown"]


def test_generate_for_ref_uses_ref_source_text(monkeypatch, tmp_path):
    from config import settings

    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("agents.doc_generation.service.get_ref_source_text", lambda ref: "")

    out = service.generate_for_ref("PROC-1", ["bsd"], overrides={})

    assert len(out) == 1
    filename, data = out[0]
    assert filename.endswith(".docx")
    assert data[:2] == b"PK"


def test_generate_for_ref_does_not_touch_thread_storage(monkeypatch):
    """Ref-based generation must not require/validate a thread_id at all."""
    monkeypatch.setattr("agents.doc_generation.service.get_ref_source_text", lambda ref: "")

    def boom(*a, **k):
        raise AssertionError("save_generated_file (thread storage) must not be called")

    monkeypatch.setattr("agents.doc_generation.service.save_generated_file", boom)

    out = service.generate_for_ref("PROC-1", ["oss"], overrides={})
    assert len(out) == 1 and out[0][1][:2] == b"PK"


def test_detect_blank_text_file_has_no_types(monkeypatch):
    docs = [{"filename": "blank.pdf", "text": "   "}]
    monkeypatch.setattr("agents.doc_generation.service.get_source_documents", lambda tid: docs)
    monkeypatch.setattr(
        "agents.doc_generation.service.classify_documents",
        lambda text: (_ for _ in ()).throw(AssertionError("should not classify blank")),
    )
    out = service.detect("any-thread")
    assert out["documents"] == [{"filename": "blank.pdf", "doc_types": []}]
    assert out["doc_types"] == []
