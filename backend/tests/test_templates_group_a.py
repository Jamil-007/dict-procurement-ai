from pathlib import Path
import openpyxl
from docxtpl import DocxTemplate

def test_xlsx_present_with_sheets():
    assert "Indicative PPMP (2)" in openpyxl.load_workbook("templates/forms/ppmp.xlsx").sheetnames
    assert "APP " in openpyxl.load_workbook("templates/forms/app.xlsx").sheetnames

def test_docx_have_expected_vars():
    for f, needed in [("market.docx", {"procuring_entity", "project_name"}),
                      ("contract.docx", {"project_title", "procuring_entity", "disclaimer"})]:
        vars_ = DocxTemplate(f"templates/forms/{f}").get_undeclared_template_variables()
        assert needed.issubset(vars_), (f, vars_)
